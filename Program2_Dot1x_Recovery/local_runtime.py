"""Local persistence and process-owned tasks. No credentials are persisted."""

from __future__ import annotations

import csv
import io
import json
import sqlite3
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path


def now():
    return datetime.now(timezone.utc).isoformat()


class Store:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS inventory(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS snapshots(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            """)

    @contextmanager
    def connect(self):
        connection = sqlite3.connect(self.path, timeout=30)
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def put(self, table, key, data):
        if table not in {"inventory", "tasks", "snapshots"}:
            raise ValueError("Unknown table")
        with self.connect() as db:
            db.execute(f"INSERT OR REPLACE INTO {table} VALUES (?, ?)", (key, json.dumps(data)))

    def rows(self, table):
        if table not in {"inventory", "tasks", "snapshots"}:
            raise ValueError("Unknown table")
        with self.connect() as db:
            return [json.loads(row[0]) for row in db.execute(f"SELECT data FROM {table} ORDER BY rowid")]

    def delete_device(self, key):
        with self.connect() as db:
            db.execute("DELETE FROM inventory WHERE id=?", (key,))

    def save_device(self, data):
        allowed = {"id", "hostname", "management_ip", "wan_ip", "vendor", "model", "site", "notes"}
        if set(data) - allowed:
            raise ValueError("Inventory contains unsupported fields; credentials are not stored")
        if not data.get("hostname", "").strip():
            raise ValueError("Hostname is required")
        data = {k: str(v).strip() for k, v in data.items()}
        data["id"] = data.get("id") or uuid.uuid4().hex
        self.put("inventory", data["id"], data)
        return data

    def import_inventory(self, content, xlsx=False):
        if xlsx:
            from openpyxl import load_workbook

            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            rows = iter(sheet.values)
            headers = next(rows)
            records = [dict(zip(headers, row)) for row in rows]
        else:
            records = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        # Validate the entire input before storing any records.
        allowed = {"id", "hostname", "management_ip", "wan_ip", "vendor", "model", "site", "notes"}
        if any(set(row) - allowed or not str(row.get("hostname") or "").strip() for row in records):
            raise ValueError("Invalid inventory headers or hostname")
        return [self.save_device({k: str(v or "") for k, v in row.items()}) for row in records]

    def export_inventory(self, xlsx=False):
        fields = ["id", "hostname", "management_ip", "wan_ip", "vendor", "model", "site", "notes"]

        def safe(value):
            value = str(value or "")
            return "'" + value if value.startswith(("=", "+", "-", "@")) else value

        rows = [[safe(row.get(k)) for k in fields] for row in self.rows("inventory")]
        if xlsx:
            from openpyxl import Workbook

            book = Workbook()
            book.active.append(fields)
            for row in rows:
                book.active.append(row)
            out = io.BytesIO()
            book.save(out)
            return out.getvalue()
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(fields)
        writer.writerows(rows)
        return out.getvalue().encode()


class Tasks:
    def __init__(self, store):
        self.store = store
        self.pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix="device-task")
        self.events = {}
        self.guard = threading.Lock()
        # A process restart cannot resume an uncertain device write.
        for task in store.rows("tasks"):
            if task["state"] in {"queued", "running"}:
                task.update(
                    state="interrupted", explanation="Process restarted; inspect readback before retrying"
                )
                store.put("tasks", task["id"], task)

    def submit(self, label, items, operation):
        key = uuid.uuid4().hex
        event = threading.Event()
        self.events[key] = event
        task = dict(id=key, label=label, state="queued", created=now(), total=len(items), results=[])
        self.store.put("tasks", key, task)

        def run():
            task["state"] = "running"
            self.store.put("tasks", key, task)
            for item in items:
                if event.is_set():
                    break
                try:
                    result = operation(item, key)
                except Exception as exc:
                    # Never persist transport exception text, which may contain credentials.
                    result = dict(
                        target=item.get("hostname", item.get("device", "")),
                        status="FAILED",
                        explanation=type(exc).__name__ + ": operation failed; inspect device before retry",
                    )
                result["timestamp"] = now()
                task["results"].append(result)
                self.store.put("tasks", key, task)
                if result.get("fatal"):
                    break
            task["state"] = (
                "cancelled"
                if event.is_set()
                else (
                    "failed"
                    if any(r.get("status") in {"FAILED", "UNKNOWN"} for r in task["results"])
                    else "completed"
                )
            )
            self.store.put("tasks", key, task)

        self.pool.submit(run)
        return key

    def cancel(self, key):
        if key in self.events:
            self.events[key].set()

    def close(self):
        for event in self.events.values():
            event.set()
        self.pool.shutdown(wait=True)
