import io

import pytest
from engine import Mapping, Recovery, checksum, normalize, port_name, vlan_id
from local_runtime import Store
from openpyxl import Workbook


@pytest.fixture
def system(tmp_path):
    store = Store(tmp_path / "state.db")
    recovery = Recovery(store)
    recovery.seed("SW1", "Gi1/0/1")
    return store, recovery


def test_priority_and_collision():
    mapping = Mapping()
    mapping.reference = {"asw_9_1": 30, "asw_9_2": 40}
    mapping.special.add(("asw_9_1", port_name("Gi1/0/1")))
    assert mapping.resolve(" ASW_9_1 ", "GigabitEthernet1/0/1", 50) == (50, "manual per-interface")
    assert mapping.resolve("ASW_9_1", "Gi1/0/1") == (27, "special user")
    assert mapping.resolve("ASW_9_2", "Gi1/0/1") == (40, "exact reference")
    with pytest.raises(ValueError):
        mapping.resolve("ASW_9_10", "Gi1/0/1")


@pytest.mark.parametrize("value", [0, 4095, -1, True, "27;reload", "27.5"])
def test_invalid_vlan(value):
    with pytest.raises((ValueError, TypeError)):
        vlan_id(value)


def test_workbook_normalization():
    book = Workbook()
    sheet = book.active
    sheet.title = "Sum_Vlan"
    sheet.append(["SW Name", "VLAN"])
    sheet.append(["SW1", 30])
    sheet = book.create_sheet("Sum_switch")
    sheet.append(["SW Name", "IP"])
    sheet.append(["SW1", "192.0.2.1"])
    sheet = book.create_sheet("Specail user")
    sheet.append(["SW Name", "Interface"])
    sheet.append(["SW1", "GE1/0/1"])
    out = io.BytesIO()
    book.save(out)
    mapping = Mapping.from_xlsx(out.getvalue())
    assert mapping.resolve("SW1", "Gi1/0/1")[0] == 27
    assert len(mapping.inventory) == 1


def test_snapshot_apply_restore_preserves_all(system):
    store, recovery = system
    key = (normalize("SW1"), port_name("Gi1/0/1"))
    before = checksum(recovery.devices[key])
    plan = recovery.plan("SW1", "Gi1/0/1", 30, "manual", "session")
    assert '-  "data_vlan": 10' in plan["diff"]
    assert checksum(recovery.devices[key]) == before
    recovery.apply(plan["id"])
    assert recovery.devices[key]["voice_vlan"] == 20
    assert recovery.devices[key]["authentication"] == []
    assert recovery.devices[key]["qos"] == "preserved"
    recovery.restore(plan["id"])
    assert checksum(recovery.devices[key]) == before
    assert store.rows("snapshots")[0]["state"] == "restored"


def test_restore_conflict_and_integrity(system):
    store, recovery = system
    plan = recovery.plan("SW1", "Gi1/0/1", 27, "special", "session")
    recovery.apply(plan["id"])
    key = (normalize("SW1"), port_name("Gi1/0/1"))
    recovery.devices[key]["description"] = "out of band change"
    with pytest.raises(ValueError, match="conflict"):
        recovery.restore(plan["id"])
    snapshot = store.rows("snapshots")[0]
    snapshot["original"]["voice_vlan"] = 99
    store.put("snapshots", plan["id"], snapshot)
    with pytest.raises(ValueError, match="integrity"):
        recovery.restore(plan["id"])


def test_snapshot_failure_prevents_write(system, monkeypatch):
    store, recovery = system
    before = str(recovery.devices)

    def fail(*args):
        raise OSError("disk full")

    monkeypatch.setattr(store, "put", fail)
    with pytest.raises(OSError):
        recovery.plan("SW1", "Gi1/0/1", 30, "manual", "session")
    assert str(recovery.devices) == before


@pytest.mark.parametrize("role", ["trunk", "management", "port-channel", "unknown"])
def test_excluded_ports(system, role):
    _, recovery = system
    recovery.devices[("sw1", port_name("Gi1/0/1"))]["type"] = role
    with pytest.raises(ValueError, match="Excluded"):
        recovery.plan("SW1", "Gi1/0/1", 30, "manual", "session")


def test_live_block_and_restart(system):
    store, recovery = system
    with pytest.raises(NotImplementedError):
        recovery.plan("SW1", "Gi1/0/1", 30, "manual", "session", live=True)
    plan = recovery.plan("SW1", "Gi1/0/1", 30, "manual", "session")
    recovery.apply(plan["id"])
    restarted = Recovery(store)
    with pytest.raises(ValueError):
        restarted.restore(plan["id"])


def test_cisco_dynamic_vlan_and_voice_parser():
    from adapters import inspect_config

    text = """interface GigabitEthernet1/0/1
 switchport mode access
 switchport voice vlan 20
 authentication port-control auto
 mab
 dot1x pae authenticator
 service-policy input VOICE
 spanning-tree portfast"""
    config = inspect_config("cisco", "GigabitEthernet1/0/1", text)
    assert config.data_vlan is None  # A RADIUS assignment is not inferred as static.
    assert config.voice_vlan == 20 and config.dynamic
    assert "service-policy input VOICE" in config.lines


def test_huawei_hybrid_auth_profile_parser():
    from adapters import inspect_config

    text = """interface GigabitEthernet0/0/1
 port link-type hybrid
 port hybrid pvid vlan 10
 port hybrid untagged vlan 10
 port hybrid tagged vlan 20
 voice-vlan 20 enable
 authentication-profile CORP
 stp edged-port enable"""
    config = inspect_config("huawei", "GigabitEthernet0/0/1", text)
    assert config.data_vlan == 10 and config.voice_vlan == 20 and config.dynamic
    assert "port hybrid tagged vlan 20" in config.lines
    with pytest.raises(ValueError, match="Ambiguous"):
        inspect_config(
            "huawei", "GigabitEthernet0/0/1", text.replace("untagged vlan 10", "untagged vlan 10 20")
        )


def test_cli_identity_trunk_and_unsupported():
    from adapters import inspect_config, require_live_adapter

    with pytest.raises(ValueError, match="identity"):
        inspect_config("cisco", "Gi1/0/2", "interface Gi1/0/1")
    with pytest.raises(ValueError, match="Excluded"):
        inspect_config("cisco", "Gi1/0/1", "interface Gi1/0/1\nswitchport mode trunk")
    with pytest.raises(NotImplementedError):
        require_live_adapter()


def test_pattern_floor_side_and_no_prefix_collision():
    book = Workbook()
    sheet = book.active
    sheet.title = "Sum Vlan"
    sheet.append(["Pattern", "Floor", "Side", "VLAN"])
    sheet.append(["ASW_9_*", 9, "left", 30])
    sheet = book.create_sheet("Sum switch")
    sheet.append(["SW Name", "Floor", "Side"])
    sheet.append(["ASW_9_1", 9, "left"])
    sheet.append(["ASW_9_2", 9, "right"])
    out = io.BytesIO()
    book.save(out)
    mapping = Mapping.from_xlsx(out.getvalue())
    assert mapping.resolve("ASW_9_1", "Gi1/0/1") == (30, "pattern reference")
    with pytest.raises(ValueError):
        mapping.resolve("ASW_9_2", "Gi1/0/1")
    with pytest.raises(ValueError):
        mapping.resolve("ASW_90_1", "Gi1/0/1")
    mapping.reference["asw_9_1"] = 40
    with pytest.raises(ValueError, match="unambiguous"):
        mapping.resolve("ASW_9_1", "Gi1/0/1")
