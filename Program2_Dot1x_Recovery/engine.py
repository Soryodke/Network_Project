"""Fail-closed recovery planning, durable snapshots, and conflict-aware restore.

Live writes deliberately require a separate firmware-validated adapter.
"""

from __future__ import annotations

import copy
import difflib
import hashlib
import io
import json
import re
import threading
import uuid

from local_runtime import now


def normalize(value):
    return re.sub(r"\s+", "", str(value or "")).casefold()


def port_name(value):
    value = str(value).strip()
    value = re.sub(r"^(?:Gi|GigabitEthernet)(?=\d)", "GigabitEthernet", value, flags=re.I)
    value = re.sub(r"^(?:GE)(?=\d)", "GigabitEthernet", value, flags=re.I)
    if not re.fullmatch(r"(?:GigabitEthernet|TenGigabitEthernet|Ethernet)\d+(?:/\d+){1,2}", value, re.I):
        raise ValueError("Only explicit physical Ethernet access ports are eligible")
    return value.casefold()


def vlan_id(value):
    if isinstance(value, bool) or str(value).strip() != str(int(value)) or not 1 <= int(value) <= 4094:
        raise ValueError("VLAN must be an integer from 1 to 4094")
    return int(value)


class Mapping:
    def __init__(self):
        self.reference, self.special, self.inventory = {}, set(), []
        self.patterns = []

    @classmethod
    def from_xlsx(cls, content):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        mapping = cls()
        sheets = {normalize(s.title).replace("_", ""): s for s in book}
        if len(sheets) != len(book.sheetnames):
            raise ValueError("Ambiguous worksheet names")
        for name in ("sumvlan", "sumswitch"):
            if name not in sheets:
                raise ValueError("Required worksheet missing: " + name)

        def records(sheet):
            rows = iter(sheet.values)
            headers = [normalize(h) for h in next(rows)]
            if len(set(headers)) != len(headers):
                raise ValueError("Duplicate workbook headers")
            return [dict(zip(headers, row)) for row in rows if any(v is not None for v in row)]

        for row in records(sheets["sumvlan"]):
            name = normalize(row.get("swname") or row.get("switchname"))
            pattern = normalize(row.get("pattern"))
            vlan = vlan_id(row.get("vlan"))
            if name and pattern:
                raise ValueError("Choose SW Name or Pattern, not both")
            if pattern:
                # A wildcard may only replace a whole underscore-delimited token.
                if any("*" in token and token != "*" for token in pattern.split("_")) or "?" in pattern:
                    raise ValueError("Pattern wildcard must occupy a whole underscore-delimited token")
                expression = (
                    "^"
                    + "_".join("[^_]+" if token == "*" else re.escape(token) for token in pattern.split("_"))
                    + "$"
                )
                mapping.patterns.append(
                    (expression, normalize(row.get("floor")), normalize(row.get("side")), vlan)
                )
            elif name:
                if name in mapping.reference:
                    raise ValueError("Duplicate or ambiguous VLAN mapping")
                mapping.reference[name] = vlan
            else:
                raise ValueError("Sum_Vlan requires SW Name or explicit Pattern and VLAN")
        special = sheets.get("specialuser", sheets.get("specailuser"))
        if special is not None:
            for row in records(special):
                mapping.special.add((normalize(row.get("swname")), port_name(row.get("interface"))))
        mapping.inventory = records(sheets["sumswitch"])
        return mapping

    def resolve(self, device, port, manual=None):
        if manual is not None:
            return vlan_id(manual), "manual per-interface"
        key = (normalize(device), port_name(port))
        if key in self.special:
            return 27, "special user"
        candidates = []
        if key[0] in self.reference:
            candidates.append((self.reference[key[0]], "exact reference"))
        metadata = [
            row for row in self.inventory if normalize(row.get("swname") or row.get("switchname")) == key[0]
        ]
        for expression, floor, side, vlan in self.patterns:
            if not re.fullmatch(expression, key[0]):
                continue
            if floor or side:
                if len(metadata) != 1:
                    raise ValueError("Floor/side mapping requires one unambiguous Sum_switch record")
                if floor and normalize(metadata[0].get("floor")) != floor:
                    continue
                if side and normalize(metadata[0].get("side")) != side:
                    continue
            candidates.append((vlan, "pattern reference"))
        if len(candidates) != 1:
            raise ValueError("No unambiguous reference VLAN")
        return candidates[0]


def checksum(config):
    return hashlib.sha256(json.dumps(config, sort_keys=True).encode()).hexdigest()


class Recovery:
    def __init__(self, store):
        self.store = store
        self.devices = {}
        self.lock = threading.RLock()

    def seed(self, device, port, vendor="cisco"):
        key = (normalize(device), port_name(port))
        with self.lock:
            self.devices.setdefault(
                key,
                dict(
                    vendor=vendor,
                    model="SIMULATED",
                    firmware="simulation-1",
                    type="access",
                    data_vlan=10,
                    voice_vlan=20,
                    authentication=["dot1x", "mab"],
                    qos="preserved",
                    stp="edge",
                    lldp=True,
                    description="Synthetic access port",
                    service_policy="VOICE",
                    extra={},
                    uplink=False,
                ),
            )
        return key

    def plan(self, device, port, vlan, source, session, live=False):
        if live:
            raise NotImplementedError("Live recovery requires a firmware-validated adapter; no commands sent")
        key = (normalize(device), port_name(port))
        with self.lock:
            original = copy.deepcopy(self.devices[key])
            if original["vendor"] not in {"cisco", "huawei"} or original["firmware"] != "simulation-1":
                raise ValueError("Unsupported adapter or firmware")
            if original["type"] != "access" or original.get("uplink") or original.get("management"):
                raise ValueError("Excluded port role")
            vlan = vlan_id(vlan)
            if vlan not in {10, 20, 27, 30, 40, 50}:
                raise ValueError("Recovery VLAN does not exist in simulated VLAN database")
            after = copy.deepcopy(original)
            after.update(data_vlan=vlan, authentication=[])
            snapshot = dict(
                id=uuid.uuid4().hex,
                device=device,
                port=port_name(port),
                session=session,
                created=now(),
                original=original,
                original_hash=checksum(original),
                recovery=after,
                recovery_hash=checksum(after),
                source=source,
                state="planned",
            )
            # Failure here prevents any device mutation.
            self.store.put("snapshots", snapshot["id"], snapshot)
            snapshot["diff"] = "\n".join(
                difflib.unified_diff(
                    json.dumps(original, indent=2, sort_keys=True).splitlines(),
                    json.dumps(after, indent=2, sort_keys=True).splitlines(),
                    fromfile="original",
                    tofile="recovery",
                )
            )
            return snapshot

    def apply(self, snapshot_id):
        with self.lock:
            snapshot = self._snapshot(snapshot_id)
            key = (normalize(snapshot["device"]), snapshot["port"])
            if snapshot["state"] != "planned" or checksum(self.devices[key]) != snapshot["original_hash"]:
                raise ValueError("Current configuration conflicts with the reviewed plan")
            self.devices[key] = copy.deepcopy(snapshot["recovery"])
            if checksum(self.devices[key]) != snapshot["recovery_hash"]:
                raise RuntimeError("Recovery readback failed")
            snapshot.update(state="recovered", recovered=now())
            self.store.put("snapshots", snapshot_id, snapshot)
            return dict(
                target=snapshot["device"],
                interface=snapshot["port"],
                status="RECOVERED",
                explanation="SIMULATED readback verified; voice/QoS/STP preserved",
            )

    def _snapshot(self, snapshot_id):
        snapshot = next(s for s in self.store.rows("snapshots") if s["id"] == snapshot_id)
        if (
            checksum(snapshot["original"]) != snapshot["original_hash"]
            or checksum(snapshot["recovery"]) != snapshot["recovery_hash"]
        ):
            raise ValueError("Snapshot integrity check failed")
        return snapshot

    def restore(self, snapshot_id):
        with self.lock:
            snapshot = self._snapshot(snapshot_id)
            key = (normalize(snapshot["device"]), snapshot["port"])
            if (
                snapshot["state"] != "recovered"
                or key not in self.devices
                or checksum(self.devices[key]) != snapshot["recovery_hash"]
            ):
                raise ValueError("Restore conflict or unknown current device state; no changes sent")
            self.devices[key] = copy.deepcopy(snapshot["original"])
            if checksum(self.devices[key]) != snapshot["original_hash"]:
                raise RuntimeError("Restore readback failed")
            snapshot.update(state="restored", restored=now())
            self.store.put("snapshots", snapshot_id, snapshot)
            return dict(
                target=snapshot["device"],
                interface=snapshot["port"],
                status="RESTORED",
                explanation="SIMULATED exact original authentication and VLAN state restored",
            )
